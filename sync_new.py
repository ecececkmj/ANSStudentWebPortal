"""
Reads the Excel masterlist and copies it into the allowed_lrns table.

Excel columns (row 1 is the header):  A = LRN   B = Full Name   C = Section

Usage: put your Excel file in the project folder and name it  masterlist.xlsx
(or leave just one .xlsx file in the folder). Edit it and SAVE - the website notices
the change by itself (start_watcher) and updates the database within a few seconds.
No restart needed. sync_masterlist() also runs once every time the website starts.
"""
import glob
import os
import re
import sqlite3
import threading
import time
from decimal import Decimal, InvalidOperation

EXCEL_NAME = 'masterlist.xlsx'

# True  = Excel is the source of truth: an LRN deleted from Excel is also removed from allowed_lrns.
# False = only add / update, never remove.
# (Safety: nothing is ever removed if the Excel could not be read or has no valid rows.
#  Students who already registered are never deleted - only the allowed_lrns list changes.)
REMOVE_MISSING = True

# Sections that are NOT Grade 11 but may use the website (the research group). They are stored as
# grade 11, so they sign up like everyone else and show in the teacher's All Students list.
EXCEPTION_SECTIONS = {'12leda'}
CHECK_EVERY = 3   # seconds between checks for a saved change


STAFF_NAME = 'authorized_staff.xlsx'


def _folders():
    """Folders to look in: where the app runs from, where this file lives, and the folder above."""
    here = os.path.dirname(os.path.abspath(__file__))
    out = []
    for f in (os.getcwd(), here, os.path.dirname(here)):
        if f not in out:
            out.append(f)
    return out


def _xlsx_files(folder):
    return [f for f in glob.glob(os.path.join(folder, '*.xlsx')) if not os.path.basename(f).startswith('~$')]


def _is_staff_file(path):
    name = os.path.basename(path).lower()
    return 'staff' in name or 'authorized' in name


def find_excel_file():
    """The STUDENT list: masterlist.xlsx, or the only .xlsx that is not the staff list."""
    for folder in _folders():
        exact = os.path.join(folder, EXCEL_NAME)
        if os.path.exists(exact):
            return exact
        others = [f for f in _xlsx_files(folder) if not _is_staff_file(f)]
        if len(others) == 1:
            return others[0]
    return None


def find_staff_file():
    """The AUTHORIZED STAFF list: authorized_staff.xlsx, or the only .xlsx with 'staff' or 'authorized' in its name."""
    for folder in _folders():
        exact = os.path.join(folder, STAFF_NAME)
        if os.path.exists(exact):
            return exact
        found = [f for f in _xlsx_files(folder) if _is_staff_file(f)]
        if len(found) == 1:
            return found[0]
    return None


def clean_lrn(value):
    """Turns 120123456789, '120123456789', 1.20123456789E+11 into '120123456789'. Returns None if not 12 digits."""
    if value is None:
        return None
    try:
        if isinstance(value, (int, float)):
            text = str(int(Decimal(str(value))))
        else:
            text = str(value).strip()
            if re.fullmatch(r'\d+(\.0+)?', text):
                text = text.split('.')[0]
            elif re.search(r'[eE]', text):
                text = str(int(Decimal(text)))
    except (InvalidOperation, ValueError):
        return None
    text = text.zfill(12) if text.isdigit() and len(text) < 12 else text
    return text if (text.isascii() and text.isdigit() and len(text) == 12) else None


def sync_masterlist(db_path='database.db'):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print('[masterlist] openpyxl is not installed. Run:  pip install openpyxl')
        return

    path = find_excel_file()
    if not path:
        print(f'[masterlist] No Excel file found. Put "{EXCEL_NAME}" in the project folder.')
        return

    try:
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(min_row=2, max_col=3, values_only=True))
    except Exception as e:
        print(f'[masterlist] Could not read {path}: {e}  (is it open in Excel? close it and restart)')
        return

    conn = sqlite3.connect(db_path)
    conn.execute('''CREATE TABLE IF NOT EXISTS allowed_lrns (
        lrn TEXT PRIMARY KEY, full_name TEXT, section TEXT, grade_level INTEGER)''')

    added = updated = skipped = 0
    seen = set()
    for lrn_raw, name, section in rows:
        if not lrn_raw and not name:
            continue                                   # blank row
        lrn = clean_lrn(lrn_raw)
        name = (str(name).strip() if name else '')
        section = (str(section).strip() if section else '')
        grade = re.match(r'\s*(\d+)', section)
        if not lrn or not name or not grade:
            print(f'[masterlist] SKIPPED a row (check LRN / name / section): {lrn_raw!r}, {name!r}, {section!r}')
            skipped += 1
            continue
        seen.add(lrn)
        grade = int(grade.group(1))
        if re.sub(r'[^a-z0-9]', '', section.lower()) in EXCEPTION_SECTIONS:
            grade = 11

        cur = conn.execute('UPDATE allowed_lrns SET full_name=?, section=?, grade_level=? WHERE lrn=?',
                           (name, section, grade, lrn))
        if cur.rowcount == 0:
            conn.execute('INSERT INTO allowed_lrns (lrn, full_name, section, grade_level) VALUES (?,?,?,?)',
                         (lrn, name, section, grade))
            added += 1
        else:
            updated += 1

    removed = 0
    if REMOVE_MISSING and seen and skipped == 0:
        # skipped == 0: if any row was unreadable, don't risk deleting a student by mistake
        for (old,) in conn.execute('SELECT lrn FROM allowed_lrns').fetchall():
            if old not in seen:
                conn.execute('DELETE FROM allowed_lrns WHERE lrn=?', (old,))
                removed += 1
    elif REMOVE_MISSING and skipped:
        print('[masterlist] Removal skipped because some rows have errors - fix them first.')

    conn.commit()
    conn.close()
    print(f'[masterlist] Read {path}: {added} added, {updated} updated, {removed} removed.')


def clean_text_id(value):
    """Employee IDs / LRNs: 120123456789, '120123456789', 1.20123456789E+11, 'EMP-001' -> text."""
    if value is None:
        return ''
    if isinstance(value, (int, float)):
        try:
            return str(int(Decimal(str(value))))
        except (InvalidOperation, ValueError):
            return str(value).strip()
    text = str(value).strip()
    if re.fullmatch(r'\d+\.0+', text):
        text = text.split('.')[0]
    return text


def clean_role(value):
    text = str(value or '').strip().lower()
    if text in ('sslg', 'ssg', 'sslg officer', 'student government', 'officer'):
        return 'sslg'
    if text in ('teacher', 'faculty', 'staff'):
        return 'teacher'
    return None


def sync_staff(db_path='database.db'):
    """Excel columns (row 1 = header):  A = Access Code   B = Full Name   C = Role (Teacher or SSLG)
    The Access Code is the person's key: it identifies them, so every person needs their OWN code."""
    conn = sqlite3.connect(db_path)
    conn.execute('''CREATE TABLE IF NOT EXISTS allowed_staff (
        staff_id TEXT PRIMARY KEY, full_name TEXT NOT NULL, role TEXT NOT NULL, access_code TEXT)''')
    conn.commit()

    try:
        from openpyxl import load_workbook
    except ImportError:
        conn.close()
        print('[staff] openpyxl is not installed. Run:  pip install openpyxl')
        return

    path = find_staff_file()
    if not path:
        conn.close()
        print(f'[staff] No staff Excel found. Put "{STAFF_NAME}" in the project folder to allow teachers / SSLG to sign up.')
        return

    try:
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(min_row=2, max_col=3, values_only=True))
    except Exception as e:
        conn.close()
        print(f'[staff] Could not read {path}: {e}  (is it open in Excel? save and close it)')
        return

    added = updated = skipped = 0
    seen = set()
    for code_raw, name, role_raw in rows:
        if not code_raw and not name:
            continue                                   # blank row
        code = clean_text_id(code_raw)
        name = str(name).strip() if name else ''
        role = clean_role(role_raw)
        if not code and name and role:
            print(f'[staff] {name} has no access code yet - run make_codes.py to create one. (Not on the list until then.)')
            continue
        if len(code) < 6 or not name or not role:
            print(f'[staff] SKIPPED a row (code must be 6+ characters, name needed, role = Teacher or SSLG): {code_raw!r}, {name!r}, {role_raw!r}')
            skipped += 1
            continue
        if code.lower() in seen:
            print(f'[staff] SKIPPED a row: the access code {code!r} is used twice. Every person needs a different code.')
            skipped += 1
            continue
        seen.add(code.lower())
        cur = conn.execute('UPDATE allowed_staff SET full_name=?, role=? WHERE staff_id=?', (name, role, code))
        if cur.rowcount == 0:
            conn.execute('INSERT INTO allowed_staff (staff_id, full_name, role, access_code) VALUES (?,?,?,NULL)', (code, name, role))
            added += 1
        else:
            updated += 1

    removed = 0
    if REMOVE_MISSING and seen and skipped == 0:
        for (old,) in conn.execute('SELECT staff_id FROM allowed_staff').fetchall():
            if old.lower() not in seen:
                conn.execute('DELETE FROM allowed_staff WHERE staff_id=?', (old,))
                removed += 1
    elif REMOVE_MISSING and skipped:
        print('[staff] Removal skipped because some rows have errors - fix them first.')

    conn.commit()
    conn.close()
    print(f'[staff] Read {path}: {added} added, {updated} updated, {removed} removed.')


_watcher_started = False


def _stamp(path):
    return (os.path.getmtime(path), os.path.getsize(path)) if path else None


def start_watcher(db_path='database.db'):
    """Starts a background thread that re-syncs the student list AND the staff list whenever an Excel file is saved."""
    global _watcher_started
    if _watcher_started:
        return
    # Flask debug mode runs the app twice; only start the watcher in the real worker.
    if os.environ.get('FLASK_DEBUG') == '1' and os.environ.get('WERKZEUG_RUN_MAIN') != 'true':
        return
    _watcher_started = True

    def loop():
        last = None
        while True:
            try:
                current = (_stamp(find_excel_file()), _stamp(find_staff_file()))
                if last is not None and current != last:
                    time.sleep(2)          # let Excel finish writing the file
                    if current[0] != last[0]:
                        sync_masterlist(db_path)
                    if current[1] != last[1]:
                        sync_staff(db_path)
                    current = (_stamp(find_excel_file()), _stamp(find_staff_file()))
                last = current
            except Exception as e:
                print(f'[masterlist] watcher error: {e}')
            time.sleep(CHECK_EVERY)

    threading.Thread(target=loop, daemon=True, name='masterlist-watcher').start()
    print(f'[masterlist] Watching the Excel files for changes (every {CHECK_EVERY}s).')


if __name__ == '__main__':
    sync_masterlist()
    sync_staff()
